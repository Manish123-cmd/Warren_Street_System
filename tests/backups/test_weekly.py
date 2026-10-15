import copy
from datetime import datetime, timedelta
from io import BytesIO
import json
import unittest

import mongomock
from openpyxl import load_workbook

from backups.core import CATEGORIES, MIME, UTC, config_sources, period_bounds, reporting_week, run_names, workbook
from backups.runner import apply_retention, backup, retention_plan
from backups.storage import FOLDER, MongoStore, checksum


class FakeDrive:
    def __init__(self):
        self.files = {}
        self.content = {}
        self.counter = 0
        self.fail = None
        self.fail_after_create = False
        self.upload_count = 0
        self.trashed = []

    def new_id(self):
        self.counter += 1
        return f'file-{self.counter}'

    def get(self, file_id):
        return copy.deepcopy(self.files.get(file_id))

    def folders(self, parent, name):
        return [copy.deepcopy(f) for f in self.files.values() if f['mimeType'] == FOLDER and f['name'] == name
                and parent in f['parents'] and not f.get('trashed')]

    def create_folder(self, file_id, parent, name):
        self.files[file_id] = {'id': file_id, 'name': name, 'parents': [parent], 'mimeType': FOLDER}

    def upload(self, file_id, parent, name, content, properties):
        category = properties['warrenCategory']
        if category == self.fail and not self.fail_after_create:
            raise TimeoutError('Fake network failure')
        if file_id in self.files:
            raise AssertionError('Duplicate create attempted')
        self.upload_count += 1
        self.files[file_id] = {'id': file_id, 'name': name, 'parents': [parent], 'mimeType': MIME,
            'md5Checksum': checksum(content), 'size': str(len(content)), 'appProperties': properties}
        self.content[file_id] = content
        if category == self.fail:
            raise TimeoutError('Upload succeeded but response was lost')

    def trash(self, file_id):
        self.files[file_id]['trashed'] = True
        self.trashed.append(file_id)


class BackupTests(unittest.TestCase):
    def setUp(self):
        self.database = mongomock.MongoClient(tz_aware=True).cafe
        self.store = MongoStore(self.database, 'backup_manifests')
        self.drive = FakeDrive()
        self.now = datetime(2026, 10, 12, 20, tzinfo=UTC)
        self.sources = {c: {'collection': f'test_{c}', 'date_field': 'date', 'date_type': 'date-string',
                           'fields': {'Product': 'product', 'Quantity': 'quantity', 'Reason': 'reason'}} for c in CATEGORIES}
        for spec in self.sources.values():
            self.database[spec['collection']].insert_many([
                {'date': '2026-10-04', 'product': 'Excluded before week', 'quantity': 99},
                {'date': '2026-10-05', 'product': 'Butter Croissant', 'quantity': 2},
                {'date': '2026-10-11', 'product': 'Cruffin', 'quantity': 1},
                {'date': '2026-10-12', 'product': 'Excluded next week', 'quantity': 88},
            ])

    def run_backup(self, now=None):
        return backup(self.store, self.drive, self.sources, 'approved-parent', now or self.now)

    def test_previous_completed_week_and_manual_recovery(self):
        self.assertEqual(tuple(map(str, reporting_week(self.now))), ('2026-10-05', '2026-10-11'))
        self.assertEqual(str(reporting_week(datetime(2026, 10, 18, tzinfo=UTC))[0]), '2026-10-05')
        self.assertEqual(str(reporting_week(self.now, '2026-09-28')[0]), '2026-09-28')
        for invalid in ('2026-10-12', '2026-10-06', '2026-10-19'):
            with self.assertRaises(ValueError):
                reporting_week(self.now, invalid)

    def test_london_timestamp_gmt_bst_and_midnight(self):
        self.assertEqual(run_names(self.now), ('2026-10-12_21-00-00', ('2026', 'October')))
        self.assertEqual(run_names(datetime(2026, 12, 7, 20, tzinfo=UTC))[0], '2026-12-07_20-00-00')
        self.assertEqual(run_names(datetime(2026, 9, 30, 23, 30, tzinfo=UTC))[1], ('2026', 'October'))

    def test_dst_reporting_boundaries(self):
        from datetime import date
        start, end = period_bounds(date(2026, 10, 19))
        self.assertEqual(end - start, timedelta(hours=169))
        start, end = period_bounds(date(2026, 3, 23))
        self.assertEqual(end - start, timedelta(hours=167))

    def test_four_workbooks_same_timestamp_metadata_and_correct_rows(self):
        result = self.run_backup()
        self.assertEqual(result['status'], 'complete')
        self.assertEqual(set(result['artifacts']), set(CATEGORIES))
        self.assertEqual(self.drive.upload_count, 4)
        for category, artifact in result['artifacts'].items():
            self.assertEqual(artifact['name'], f'2026-10-12_21-00-00_{category}.xlsx')
            self.assertNotIn(':', artifact['name'])
            book = load_workbook(BytesIO(self.drive.content[artifact['id']]))
            metadata = dict(book['Reporting Period'].values)
            self.assertEqual(metadata['Reporting period start (London)'], '2026-10-05')
            self.assertEqual(metadata['Reporting period end (London, inclusive)'], '2026-10-11')
            self.assertEqual(metadata['Record count'], 2)
            self.assertEqual(book.worksheets[1].max_row, 3)
            self.assertIsNone(book.worksheets[1]['D2'].value)  # Missing reason isn't fabricated.
        self.assertEqual([f['name'] for f in self.drive.files.values() if f['mimeType'] == FOLDER],
                         ['Warren Street Backups', '2026', 'October'])

    def test_repeat_reuses_verified_files(self):
        first = self.run_backup()
        self.run_backup(self.now + timedelta(days=2))
        self.assertEqual(self.drive.upload_count, 4)
        final = self.database.backup_manifests.find_one({'_id': first['_id']})
        self.assertEqual(first['timestamp'], final['timestamp'])
        self.assertEqual(first['artifacts'], final['artifacts'])

    def test_partial_failure_preserves_files_and_resumes_same_ids_and_time(self):
        self.drive.fail = 'Sales'
        with self.assertRaises(TimeoutError):
            self.run_backup()
        partial = self.database.backup_manifests.find_one({'_id': 'weekly:2026-10-05'})
        first_ids = {c: a['id'] for c, a in partial['artifacts'].items()}
        self.assertEqual(partial['status'], 'failed')
        self.assertTrue(partial['artifacts']['Stock']['verified'])
        self.drive.fail = None
        result = self.run_backup(self.now + timedelta(days=1))
        self.assertEqual(result['timestamp'], partial['timestamp'])
        self.assertEqual(self.drive.upload_count, 4)
        self.assertTrue(all(result['artifacts'][c]['id'] == file_id for c, file_id in first_ids.items()))

    def test_lost_response_after_upload_does_not_duplicate(self):
        self.drive.fail, self.drive.fail_after_create = 'Wastage', True
        with self.assertRaises(TimeoutError):
            self.run_backup()
        self.drive.fail = None
        result = self.run_backup()
        self.assertEqual(result['status'], 'complete')
        self.assertEqual(self.drive.upload_count, 4)

    def test_changed_source_during_unfinished_upload_stops_without_replacing_success(self):
        self.drive.fail = 'Sales'
        with self.assertRaises(TimeoutError):
            self.run_backup()
        self.database[self.sources['Sales']['collection']].update_one({'date': '2026-10-05'}, {'$set': {'quantity': 123}})
        self.drive.fail = None
        with self.assertRaises(RuntimeError):
            self.run_backup()
        self.assertEqual(self.drive.upload_count, 2)
        self.assertEqual(self.drive.trashed, [])

    def test_retry_across_month_keeps_initial_folder_and_timestamp(self):
        first_time = datetime(2026, 10, 31, 23, 59, tzinfo=UTC)
        first = backup(self.store, self.drive, self.sources, 'approved-parent', first_time, '2026-10-05')
        retried = backup(self.store, self.drive, self.sources, 'approved-parent', first_time + timedelta(minutes=2), '2026-10-05')
        self.assertEqual(retried['folder_path'], ['2026', 'October'])
        self.assertEqual(first['timestamp'], retried['timestamp'])
        self.assertEqual(self.drive.upload_count, 4)

    def test_deterministic_workbook_and_oversized_cells_fail_instead_of_truncating(self):
        manifest = self.run_backup()
        documents = [{'_id': 1, 'product': 'Same data'}]
        first = workbook('Stock', documents, self.sources['Stock'], manifest)
        self.assertEqual(first, workbook('Stock', documents, self.sources['Stock'], manifest))
        with self.assertRaises(ValueError):
            workbook('Stock', [{'_id': 1, 'product': 'x' * 32768}], self.sources['Stock'], manifest)

    def test_missing_source_collection_fails_and_records_failure(self):
        self.database.drop_collection(self.sources['Stock']['collection'])
        with self.assertRaises(ValueError):
            self.run_backup()
        self.assertEqual(self.database.backup_manifests.find_one({'_id': 'weekly:2026-10-05'})['status'], 'failed')
        self.assertEqual(self.drive.upload_count, 0)

    def test_corrupt_or_missing_successful_files_fail_without_overwrite(self):
        result = self.run_backup()
        artifact = result['artifacts']['Stock']
        self.drive.files[artifact['id']]['md5Checksum'] = 'corrupt'
        with self.assertRaises(RuntimeError):
            self.run_backup()
        self.assertEqual(self.drive.upload_count, 4)
        self.assertEqual(self.drive.trashed, [])

    def test_no_recipe_exports_or_implicit_collection_names(self):
        self.assertEqual(config_sources(json.dumps(self.sources), 'backup_manifests'), self.sources)
        for collection in ('recipes', 'backup_manifests'):
            invalid = copy.deepcopy(self.sources)
            invalid['Stock']['collection'] = collection
            with self.assertRaises(ValueError):
                config_sources(json.dumps(invalid), 'backup_manifests')
        with self.assertRaises(ValueError):
            config_sources('{}', 'backup_manifests')

    def test_excel_text_is_never_an_executable_formula(self):
        result = self.run_backup()
        data = workbook('Stock', [{'_id': 1, 'product': '=HYPERLINK("bad")'}], self.sources['Stock'], result)
        cell = load_workbook(BytesIO(data)).worksheets[1]['B2']
        self.assertEqual(cell.value, '=HYPERLINK("bad")')
        self.assertEqual(cell.data_type, 's')

    def test_no_live_source_mutations(self):
        before = {c: list(self.database[spec['collection']].find()) for c, spec in self.sources.items()}
        self.run_backup()
        self.assertEqual(before, {c: list(self.database[spec['collection']].find()) for c, spec in self.sources.items()})

    def test_overlapping_run_is_rejected(self):
        with self.store.lock():
            with self.assertRaises(RuntimeError):
                with MongoStore(self.database, 'backup_manifests').lock():
                    self.fail('Second runner acquired the lease')

    def test_retention_disabled_and_plan_only_never_deletes(self):
        self.run_backup()
        later = self.now + timedelta(days=100)
        self.assertEqual(retention_plan(self.store, 0, later, 'approved-parent')['backups'], [])
        self.assertEqual(len(retention_plan(self.store, 30, later, 'approved-parent')['backups']), 1)
        self.assertEqual(self.drive.trashed, [])
        with self.assertRaises(ValueError):
            apply_retention(self.store, self.drive, {}, 0, later, 'approved-parent')

    def test_explicit_retention_only_trashes_reviewed_verified_files(self):
        result = self.run_backup()
        later = self.now + timedelta(days=100)
        plan = retention_plan(self.store, 30, later, 'approved-parent')
        bad = copy.deepcopy(plan)
        bad['backups'][0]['files']['Stock'] = 'unrelated-file'
        with self.assertRaises(ValueError):
            apply_retention(self.store, self.drive, bad, 30, later, 'approved-parent')
        self.assertEqual(self.drive.trashed, [])
        self.assertEqual(apply_retention(self.store, self.drive, plan, 30, later, 'approved-parent'), 1)
        self.assertEqual(set(self.drive.trashed), {a['id'] for a in result['artifacts'].values()})

    def test_datetime_query_obeys_local_dst_boundaries(self):
        from datetime import date
        spec = {'collection': 'utc_events', 'date_field': 'at', 'date_type': 'datetime', 'fields': {'When': 'at'}}
        self.database.utc_events.insert_many([
            {'at': datetime(2026, 10, 18, 23, tzinfo=UTC)},
            {'at': datetime(2026, 10, 25, 23, 59, tzinfo=UTC)},
            {'at': datetime(2026, 10, 26, tzinfo=UTC)},
        ])
        with self.store.lock():
            self.assertEqual(len(self.store.documents(spec, date(2026, 10, 19))), 2)


if __name__ == '__main__':
    unittest.main()
