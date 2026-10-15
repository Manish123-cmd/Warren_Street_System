"""London reporting periods and four independent, readable Excel exports."""
from calendar import month_name
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
import json
import re
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from zoneinfo import ZoneInfo

from bson import json_util
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

LONDON = ZoneInfo('Europe/London')
UTC = timezone.utc
CATEGORIES = ('Stock', 'Wastage', 'Sales', 'Delivery_Notes')
MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def reporting_week(now, requested=None):
    if now.tzinfo is None:
        raise ValueError('An aware backup time is required')
    today = now.astimezone(LONDON).date()
    monday = today - timedelta(days=today.weekday())
    start = date.fromisoformat(requested) if requested else monday - timedelta(days=7)
    if start.weekday() != 0 or start >= monday:
        raise ValueError('Select a previous completed Monday-Sunday week')
    return start, start + timedelta(days=6)


def period_bounds(start):
    # Local midnights, not seven times 24 UTC hours: DST weeks differ in length.
    return tuple(datetime.combine(d, time.min, LONDON).astimezone(UTC)
                 for d in (start, start + timedelta(days=7)))


def run_names(now):
    local = now.astimezone(LONDON)
    return local.strftime('%Y-%m-%d_%H-%M-%S'), (str(local.year), month_name[local.month])


def config_sources(raw, manifest_collection):
    sources = json.loads(raw)
    if set(sources) != set(CATEGORIES):
        raise ValueError('Configure exactly Stock, Wastage, Sales and Delivery_Notes')
    for spec in sources.values():
        if (not isinstance(spec, dict) or not spec.get('collection')
                or spec['collection'] == manifest_collection
                or 'recipe' in spec['collection'].lower()
                or not isinstance(spec.get('date_field'), str) or not spec['date_field']
                or spec.get('date_type') not in ('date-string', 'datetime')
                or not isinstance(spec.get('fields'), dict) or not spec['fields']
                or any(not isinstance(k, str) or not k or not isinstance(v, str) or not v
                       for k, v in spec['fields'].items())):
            raise ValueError('Each category needs an explicit non-recipe collection, date field/type and column mapping')
    return sources


def dotted(document, path):
    value = document
    for part in path.split('.'):
        if not isinstance(value, dict) or part not in value:
            return None
        value = value[part]
    return value


def excel_value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, datetime):
        return value.astimezone(LONDON).isoformat() if value.tzinfo else value.isoformat()
    if isinstance(value, (dict, list)):
        value = json_util.dumps(value, ensure_ascii=False)
    else:
        value = str(value)
    if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
        raise ValueError('A field exceeds Excel cell limits; export rejected rather than truncated')
    return value


def write_row(sheet, values):
    sheet.append([excel_value(value) for value in values])
    # Treat supplier names, notes, etc. as text, never Excel formulas.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def workbook(category, documents, spec, manifest):
    book = Workbook()
    metadata = book.active
    metadata.title = 'Reporting Period'
    for values in (
        ('Category', category),
        ('Reporting period start (London)', manifest['week_start']),
        ('Reporting period end (London, inclusive)', manifest['week_end']),
        ('Backup run time (Europe/London)', manifest['backup_time']),
        ('Backup timestamp', manifest['timestamp']),
        ('Timezone', 'Europe/London'),
        ('Source collection', spec['collection']),
        ('Date field', spec['date_field']),
        ('Record count', len(documents)),
        ('Missing values', 'Blank means not recorded; values are not inferred.'),
        ('Scope', 'Reporting export, not a full database restore archive. Recipes excluded.'),
    ):
        write_row(metadata, values)
    sheet = book.create_sheet(category.replace('_', ' '))
    write_row(sheet, ['Source record ID', *spec['fields']])
    for document in documents:
        write_row(sheet, [document.get('_id'), *(dotted(document, path) for path in spec['fields'].values())])
    if not documents:
        write_row(metadata, ('Data status', 'No source records found in this reporting period'))
    for tab in book:
        tab.freeze_panes = 'A2'
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2759AA')
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = min(65, max(20, max(len(str(c.value or '')) for c in column) + 2))
    result = BytesIO()
    book.save(result)
    # Stable bytes allow an interrupted, not-yet-visible upload to be rebuilt safely.
    # openpyxl/ZIP otherwise embeds the current save time on every retry.
    stable = BytesIO()
    instant = datetime.fromisoformat(manifest['backup_time']).astimezone(UTC)
    stamp = instant.strftime('%Y-%m-%dT%H:%M:%SZ')
    with ZipFile(result) as original, ZipFile(stable, 'w', ZIP_DEFLATED) as archive:
        for filename in original.namelist():
            payload = original.read(filename)
            if filename == 'docProps/core.xml':
                text = payload.decode('utf-8')
                text = re.sub(r'(<dcterms:(?:created|modified)\b[^>]*>).*?(</dcterms:(?:created|modified)>)',
                              lambda match: match[1] + stamp + match[2], text)
                payload = text.encode('utf-8')
            info = ZipInfo(filename, date_time=(max(1980, instant.year), instant.month, instant.day,
                                                instant.hour, instant.minute, instant.second))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, payload)
    return stable.getvalue()
